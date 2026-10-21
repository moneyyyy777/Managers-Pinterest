import sys
import openpyxl

def extract_table(md_path):
    rows = []
    with open(md_path, 'r', encoding='utf-8') as f:
        in_table = False
        for line in f:
            if line.startswith('| № |') or line.startswith('| No |'):
                in_table = True
                continue
            if line.startswith('|---'):
                continue
            if in_table and line.startswith('|'):
                parts = [p.strip() for p in line.split('|')]
                if len(parts) >= 5:
                    title = parts[2]
                    desc = parts[3]
                    cta = parts[4]
                    if title and title.lower() != 'title':
                        rows.append({"title": title, "desc": desc, "cta": cta})
            elif in_table and not line.strip():
                break
    return rows

if __name__ == "__main__":
    if len(sys.argv) != 3:
        print("Usage: python3 process_and_save.py <response_md_file> <target_xlsx_file>")
        sys.exit(1)
        
    md_file = sys.argv[1]
    xlsx_file = sys.argv[2]
    
    rows = extract_table(md_file)
    
    if len(rows) != 100:
        print(f"ERROR: Expected 100 rows, got {len(rows)}.")
        sys.exit(1)
        
    print(f"[*] Validated 100 rows. Writing to {xlsx_file}...")
    
    wb = openpyxl.load_workbook(xlsx_file)
    sheet = wb.active
    
    # Write to L(12)=desc, M(13)=cta, N(14)=title starting at row 2
    for i, row in enumerate(rows):
        excel_row = i + 2
        sheet.cell(row=excel_row, column=12).value = row["desc"]
        sheet.cell(row=excel_row, column=13).value = row["cta"]
        sheet.cell(row=excel_row, column=14).value = row["title"]
        
    wb.save(xlsx_file)
    print(f"[*] Successfully saved to Excel!")
